"""Exports: the samples currently shown (after search and filters) as CSV, Excel or PDF."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any
from xml.sax.saxutils import escape

from . import storage

COLUMNS = [
    ("Project", "project_name"),
    ("Pen #", "pen"),
    ("Sample #", "sample_no"),
    ("Label", "label"),
    ("Pig ID #", "pig_id"),
    ("Trial day", "trial_day"),
    ("AI score", "ai_score"),
    ("Final score", "final_score"),
    ("Edited", "edited"),
    ("Description", "description"),
    ("Captured", "captured_at"),
    ("Modified", "modified_at"),
    ("Photo link", "image_url"),
    ("Crop link", "crop_url"),
]


def local_time(iso: str | None) -> str:
    """'2026-10-28T20:14:00+00:00' -> '2026-10-28 15:14' (in this computer's time zone)."""
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone().strftime("%Y-%m-%d %H:%M")
    except ValueError:
        return iso


def _row(s: dict[str, Any], base_url: str) -> list[Any]:
    values = []
    for _, key in COLUMNS:
        v = s.get(key)
        if key in ("captured_at", "modified_at"):
            v = local_time(v)
        elif key == "edited":
            v = "Yes" if v else ""
        elif key in ("image_url", "crop_url"):
            v = base_url + v if v else ""
        elif v is None:
            v = ""
        values.append(v)
    return values


def to_csv(samples: list[dict[str, Any]], base_url: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c[0] for c in COLUMNS])
    for s in samples:
        writer.writerow(_row(s, base_url))
    # utf-8-sig so Excel opens it with the right characters
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(samples: list[dict[str, Any]], base_url: str, title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Samples"
    ws.append([c[0] for c in COLUMNS])
    header_fill = PatternFill("solid", fgColor="C8102E")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    link_cols = {i + 1 for i, (_, key) in enumerate(COLUMNS) if key in ("image_url", "crop_url")}
    for s in samples:
        ws.append(_row(s, base_url))
        r = ws.max_row
        for col in link_cols:
            cell = ws.cell(row=r, column=col)
            if cell.value:
                cell.hyperlink = cell.value
                cell.value = "Open photo" if COLUMNS[col - 1][1] == "image_url" else "Open crop"
                cell.font = Font(color="0563C1", underline="single")

    widths = [22, 8, 10, 14, 10, 10, 9, 11, 8, 36, 17, 17, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("About")
    info.append(["FecalScore export"])
    info.append(["Showing", title])
    info.append(["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")])
    info.append(["Samples", len(samples)])
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 12
    info.column_dimensions["B"].width = 60

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(samples: list[dict[str, Any]], title: str, show_project: bool) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    cardinal = colors.HexColor("#C8102E")
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=8.5, leading=10.5)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=f"FecalScore — {title}",
    )

    head = ["Photo", "Sample"] + (["Project"] if show_project else []) + [
        "Pig ID #",
        "Trial day",
        "AI",
        "Score",
        "Description",
        "Captured",
    ]
    rows: list[list[Any]] = [head]
    for s in samples:
        thumb = storage.thumbnail_file(s.get("crop_path") or s.get("image_path"))
        if thumb:
            img = Image(str(thumb))
            ratio = img.imageWidth / img.imageHeight
            img.drawHeight = 0.9 * inch
            img.drawWidth = min(0.9 * inch * ratio, 1.1 * inch)
            img.drawHeight = img.drawWidth / ratio
            photo: Any = img
        else:
            photo = ""
        score = str(s["final_score"]) + (" (edited)" if s.get("edited") else "")
        row = [photo, Paragraph(f"<b>{escape(s['label'])}</b>", small)]
        if show_project:
            row.append(Paragraph(escape(s.get("project_name", "")), small))
        row += [
            s["pig_id"],
            str(s["trial_day"]),
            "" if s.get("ai_score") is None else str(s["ai_score"]),
            score,
            Paragraph(escape(s.get("description") or ""), small),
            local_time(s.get("captured_at")),
        ]
        rows.append(row)

    widths = [1.2, 0.95] + ([1.4] if show_project else []) + [0.75, 0.7, 0.4, 0.8, 0, 1.25]
    used = sum(widths)
    widths[widths.index(0)] = (10.0 - used)  # description takes the remaining width
    table = Table(rows, colWidths=[w * inch for w in widths], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), cardinal),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F5F2")]),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E2DDD7")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    story = [
        Paragraph(f"<font color='#C8102E'><b>FecalScore</b></font> — {escape(title)}", styles["Title"]),
        Paragraph(
            f"{len(samples)} sample{'s' if len(samples) != 1 else ''} · exported {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        table if samples else Paragraph("No samples match.", styles["BodyText"]),
    ]
    doc.build(story)
    return buf.getvalue()
